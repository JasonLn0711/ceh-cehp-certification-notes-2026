#!/usr/bin/env python3
"""Rebuild the M10 editable source layers from reviewed authored inputs.
Python 3.10+, python-pptx and Pillow. No TTS, sample execution, or network access.
"""
from pathlib import Path
import json,re,hashlib,base64,shutil,wave,html,math
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from PIL import Image
R=Path(__file__).resolve().parents[1]
def write(name,text):
 p=R/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding='utf-8')
def dump(name,obj):write(name,json.dumps(obj,ensure_ascii=False,indent=2)+'\n')
def words(s):return len(re.findall(r"\b\w+(?:['’-]\w+)*\b",s))
chapters=json.loads((R/'tooling/authored_chapters.json').read_text())
records={}
for file in sorted((R/'tooling').glob('authored_scenes*.json')):
 for row in json.loads(file.read_text()):
  assert row['id'] not in records,row['id']
  records[row['id']]=dict(title=row['title'],labels=row['labels'],cue=row['cue'],kind=row['kind'],narration=row['narration'],motion=row['motion'])
qs=json.loads((R/'tooling/authored_questions.json').read_text())
for i in range(50,58):
 subset=[q for q in qs if q['scene_id']==f'S{i:02}'];assert len(subset)==1
 records[f'S{i:02}']={'title':f'Knowledge check: {subset[0]["level"]}', 'labels':[subset[0]['id'],'Read all choices','One-second hold','Explain the answer'],'cue':'Read all choices before the one-second silent hold.','kind':'quiz','narration':''}
assert len(records)==58
scenes=[];ch=[]
for ci,(title,a,b,goal) in enumerate(chapters,1):
 ids=[]
 for i in range(a,b+1):
  sid=f'S{i:02}';d=records[sid];ids.append(sid);segs=[]
  if d['kind']=='quiz':
   for q in [q for q in qs if q['scene_id']==sid]:
    spoken=f"Question {int(q['id'][1:])}. {q['prompt']} "+' '.join(f"{c['letter']}. {c['text']}" for c in q['choices'])+' Choose the best answer.'
    segs.extend([dict(kind='spoken',text=spoken,question_id=q['id'],role='question_and_all_choices'),dict(kind='silence',seconds=1,question_id=q['id'],pcm_frames=48000),dict(kind='spoken',text=f"The correct answer is {q['correct_letter']}. {q['explanation']}",question_id=q['id'],role='answer_and_reasoning')])
   d['narration']=' '.join(x['text'] for x in segs if x['kind']=='spoken')
  else:segs=[dict(kind='spoken',text=d['narration'])]
  n=words(d['narration']);estimate=round(n/182*60+sum(x.get('seconds',0) for x in segs)+8,3)
  scenes.append(dict(id=sid,chapter_id=f'C{ci:02}',chapter_title=title,title=d['title'],tags=['M10','Denial of Service',d['kind']],visual_direction=d['cue'],on_screen_copy=d['labels'],narration_en_US=d['narration'],audio_segments=segs,estimated_spoken_words=n,editorial_estimate_seconds=estimate,notes='Source board; final motion and measured speech timing are the next production layer.',review_requirements=['Evidence limits remain explicit','Readable labels and caption clearance','Narration-only speed 1.5; video speed 1.0']))
 ch.append(dict(id=f'C{ci:02}',title=title,learning_goal=goal,scene_ids=ids,scene_count=len(ids),editorial_estimate_seconds=round(sum(s['editorial_estimate_seconds'] for s in scenes if s['chapter_id']==f'C{ci:02}'),3),editing_rule='Derive final duration from actual speech, animation and reading needs; no runtime ceiling.',narration_words=sum(s['estimated_spoken_words'] for s in scenes if s['chapter_id']==f'C{ci:02}')))
meta=dict(title='CEH Module 10: Denial of Service — Northstar Evidence Lab',version='1.0',status='source_authored_pending_QA',actual_final_mp4_rendered=False,language='en-US',spoken_and_visual_language='en-US',playback_rate='Only narration at 1.5x relative to M04 v3; video and animation remain 1.0.',narration_speed_multiplier=1.5,voice=dict(engine='Kokoro',voice='af_heart',baseline_speed=0.73,postprocess='exactly one atempo=1.5 on spoken clips'),base_wpm_range=[115,130],delivery_effective_wpm_range=[172.5,195],silent_think_seconds_each=1,question_count=8,scene_count=58,estimated_spoken_words=sum(s['estimated_spoken_words'] for s in scenes),duration_policy='No maximum. Estimates are planning values; actual speech, reading and motion determine final timing.',safety='SIMULATED OUTPUT: fictional offline educational evidence. Read-only synthetic CSV/JSON/PCAP incident analysis. No traffic generation, packet replay or live control deployment.',scenes=scenes)
meta['timing_estimate_basis']='Inherited 182 WPM planning assumption; not measured M10 audio.'
meta['estimated_raw_narration_seconds_at_182wpm']=round(meta['estimated_spoken_words']/182*60,3)
dump('scenes.json',meta);dump('chapters.json',ch);dump('questions.json',dict(language='en-US',namespace='M10-film-v1',question_count=8,silent_think_seconds_each=1,questions=qs))
write('SCRIPT_en-US.md','# CEH M10: Denial of Service — complete narration\n\nCanonical spoken words: '+str(meta['estimated_spoken_words'])+'. Count: Unicode word tokens with internal apostrophes/hyphens. Stage directions are unspoken.\n\n'+'\n\n'.join(f"## {s['id']} · {s['title']}\n\n"+'\n\n'.join(x['text'] if x['kind']=='spoken' else f"[TRUE SILENCE: 1.000 SECOND · {x['question_id']} · 48,000 stereo audio frames]" for x in s['audio_segments']) for s in scenes)+'\n')
write('COURSE_STRUCTURE.md','# Course structure\n\n| Chapter | Scenes | Goal | Spoken words |\n| --- | --- | --- | --- |\n'+'\n'.join(f"| {c['id']} · {c['title']} | {c['scene_ids'][0]}–{c['scene_ids'][-1]} | {c['learning_goal']} | {c['narration_words']} |" for c in ch)+'\n')
visual=[]
for s in scenes:
 d=records[s['id']];visual.append(dict(id=s['id'],chapter_id=s['chapter_id'],title=s['title'],screen_copy_exact=s['on_screen_copy'],layout=d['kind'],focal_claim=d['cue'],board=f"boards/{s['id']}.svg",caption_safe_zone=dict(x=135,y=900,width=1650,height=100),assets=['assets/mascot/tally_point.svg','STYLE_TOKENS.json','network_map.json'],animation_beats=[dict(phase='orient',action='Establish the named objects and the Northstar anchor.'),dict(phase='explain',action='Reveal each object and its labeled relationship in narration order.'),dict(phase='reason',action=d['cue']),dict(phase='hold',action='Hold for reading; derive duration from actual voice and audience need.')],continuity='Same Tally artwork, palette, campus device IDs, and source-before-interpretation logic.',timing_status='editorial_estimate_only'))
dump('visual_scenes.json',dict(version='1.0',style='M04 v3 inherited through M06 v1.1',reference='reference/HYPERFRAMES_REFERENCE_STYLE_BOARD.jpg',source_scenes='scenes.json',scene_count=58,scenes=visual))
write('STORYBOARD.md','# Editorial storyboard: all 58 scenes\n\nThese are editable design boards with explicit motion plans. Four executable animation examples accompany them. Full-film scene animation is a later production task.\n\n'+'\n\n'.join(f"## {v['id']} · {v['title']}\n\n**{v['layout']}** — {v['focal_claim']}\n\nVisible labels: {' / '.join(v['screen_copy_exact'])}.\n\nMotion: orient → reveal relationships → explain the evidence limit → reading hold. Preserve the lower caption zone. For quiz scenes, show one complete stem and A–D set at a time, read all choices, hold exactly 1.000 second with zero PCM, then reveal the answer. Each source quiz board shows one complete question for editorial review.\n\n[Editable SVG]({v['board']}) · [Narration](SCRIPT_en-US.md)" for v in visual)+'\n')
# A shared drawing model drives editable SVG and native PowerPoint shapes/text.
C=json.loads((R/'STYLE_TOKENS.json').read_text())['colors'];PAPER=C['paper'];INK=C['ink'];ACC=C['terracotta'];WHITE=C['paper_white'];GOLD=C['mustard']
mascot=R/'assets/mascot/tally_point.png';mascot_uri='data:image/png;base64,'+base64.b64encode(mascot.read_bytes()).decode()
import textwrap
def wrap(text,maxchars):return textwrap.wrap(text,width=maxchars,break_long_words=False,break_on_hyphens=False) or ['']
def model(s):
 d=records[s['id']];items=[]
 def text(x,y,w,h,t,size=38,bold=False,color=INK):items.append(dict(type='text',x=x,y=y,w=w,h=h,text=t,size=size,bold=bold,color=color))
 def box(x,y,w,h,label,fill=WHITE):items.append(dict(type='box',x=x,y=y,w=w,h=h,fill=fill));text(x+22,y+22,w-44,h-44,label,38,True)
 def line(x,y,x2,y2):items.append(dict(type='line',x=x,y=y,x2=x2,y2=y2))
 text(138,86,1500,40,f"CEH M10  /  {s['chapter_id']}  /  {s['id']}",26,True,INK)
 text(138,150,1585,160,s['title'],64,True)
 text(138,325,1530,95,d['cue'],36)
 labels=d['labels'];kind=d['kind']
 if kind=='quiz':
  subset=[q for q in qs if q['scene_id']==s['id']]
  for j,q in enumerate(subset):
   x=140;items.append(dict(type='box',x=x,y=420,w=1460,h=430,fill=WHITE))
   text(x+22,445,1416,120,q['id']+'. '+q['prompt'],32,True)
   text(x+22,595,1416,240,'\n'.join(c['letter']+'. '+c['text'] for c in q['choices']),30)
 elif kind in ['flow','timeline']:
  for j,l in enumerate(labels):
   x=140+j*400;box(x,515,340,180,l,C['apricot'] if j==0 else WHITE)
   if j<3:line(x+345,605,x+389,605)
  text(140,735,1480,90,' → '.join(labels),28,False,INK)
 elif kind in ['comparison','checklist']:
  for j,l in enumerate(labels):
   x=140+(j%2)*680;y=460+(j//2)*180;box(x,y,610,135,l,C['apricot'] if j==0 else WHITE)
 elif kind=='layers':
  for j,l in enumerate(labels):
   x=140+j*65;y=450+j*90;box(x,y,1120-j*65,82,l,C['apricot'] if j==0 else WHITE)
 elif kind=='tree':
  box(650,450,530,110,labels[0]);line(915,565,915,645)
  box(650,655,530,125,labels[1],C['apricot'])
  box(140,655,430,125,labels[2]);box(1250,655,350,125,labels[3])
  text(140,605,430,40,'Recorded context',26,False,INK);text(1250,605,350,40,'Evidence limit / context',26,False,INK)
 else:
  points=[(160,455),(860,455),(160,680),(860,680)]
  for j,((x,y),l) in enumerate(zip(points,labels)):box(x,y,540,125,l,C['apricot'] if j==0 else WHITE)
  line(710,515,848,515);line(430,590,430,670);line(1130,590,1130,670)
 if kind!='quiz':items.append(dict(type='image',x=1640,y=730,w=140,h=140,path=str(mascot)))
 text(138,858,1330,32,'Northstar Learning Center  ·  fictional educational evidence',24,False,INK)
 return items

def svg(s):
 it=model(s);out=[f'<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080"><title>{escape(s["id"]+" "+s["title"])}</title><desc>Editable source board. Caption area y=900–1000. Full animation remains to be produced.</desc><rect width="1920" height="1080" fill="{PAPER}"/>']
 out.append('<defs><style>@font-face{font-family:"Liberation Sans";src:url("../assets/fonts/LiberationSans-Regular.ttf")}@font-face{font-family:"Liberation Sans";src:url("../assets/fonts/LiberationSans-Bold.ttf");font-weight:700}</style></defs>')
 # Subtle deterministic paper edge marks; no filter-dependent noise.
 for j in range(9):out.append(f'<path d="M 65 {120+j*90} l 18 {2 if j%2 else -2}" stroke="{C["gray"]}" stroke-width="2"/>')
 for n,a in enumerate(it):
  typ=a['type'];ident=f'node-{n}'
  if typ=='box':out.append(f'<path id="{ident}" d="M {a["x"]} {a["y"]+3} l {a["w"]} -3 l 2 {a["h"]} l -{a["w"]+3} 2 Z" fill="{a["fill"]}" stroke="{INK}" stroke-width="4" stroke-linejoin="round"/>')
  elif typ=='line':
   x,y,x2,y2=a['x'],a['y'],a['x2'],a['y2'];out.append(f'<path id="{ident}" d="M {x} {y} L {x2} {y2}" fill="none" stroke="{ACC}" stroke-width="5" stroke-linecap="round"/>')
  elif typ=='image':out.append(f'<image id="{ident}" x="{a["x"]}" y="{a["y"]}" width="{a["w"]}" height="{a["h"]}" href="{mascot_uri}"/>')
  else:
   size=a['size'];maxchars=max(8,int(a['w']/(size*.53)));lines=[]
   for line in a['text'].split('\n'):lines.extend(wrap(line,maxchars))
   out.append(f'<text id="{ident}" font-family="Liberation Sans, Arial, sans-serif" font-size="{size}" font-weight="{700 if a["bold"] else 400}" fill="{a["color"]}">')
   for j,line in enumerate(lines):out.append(f'<tspan x="{a["x"]}" y="{a["y"]+size+j*size*1.18}">{escape(line)}</tspan>')
   out.append('</text>')
 out.append('</svg>');return ''.join(out)
for s in scenes:write(f'boards/{s["id"]}.svg',svg(s))
# Native editable PowerPoint: labels, boxes, connectors remain separate objects.
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
def inch(v):return Inches(v/120)
def rgb(h):return RGBColor.from_string(h.lstrip('#'))
selected=[1,4,8,12,15,18,20,24,27,29,31,34,38,43,44,47,54,58]
for index in selected:
 s=scenes[index-1];slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(PAPER)
 for a in model(s):
  typ=a['type']
  if typ=='image':slide.shapes.add_picture(a['path'],inch(a['x']),inch(a['y']),inch(a['w']),inch(a['h']));continue
  if typ=='line':
   sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,inch(a['x']),inch(a['y']),inch(a['x2']),inch(a['y2']));sh.line.color.rgb=rgb(ACC);sh.line.width=Pt(2.5);continue
  if typ=='box':
   sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,inch(a['x']),inch(a['y']),inch(a['w']),inch(a['h']));sh.fill.solid();sh.fill.fore_color.rgb=rgb(a['fill']);sh.line.color.rgb=rgb(INK);sh.line.width=Pt(2);continue
  sh=slide.shapes.add_textbox(inch(a['x']),inch(a['y']),inch(a['w']),inch(a['h']));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_top=0;tf.margin_right=0;tf.margin_bottom=0
  for j,line in enumerate(a['text'].split('\n')):
   p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='Liberation Sans';p.font.size=Pt(a['size']*.6);p.font.bold=a['bold'];p.font.color.rgb=rgb(a['color']);p.space_after=Pt(2)
 slide.notes_slide.notes_text_frame.text=f"{s['id']} — editable visual demonstration, not a rendered film frame.\n{s['narration_en_US']}\nMotion: {s['visual_direction']}"
(R/'slides').mkdir(exist_ok=True);prs.save(R/'slides/CEH_M10_Editable_Visual_Lookbook_18slides.pptx');dump('slides/selection.json',selected)
# Real reference silence, not eight files with synthetic narration attached.
p=R/'production/quiz_silence_1s_48k_stereo.wav';p.parent.mkdir(exist_ok=True)
with wave.open(str(p),'wb') as w:w.setnchannels(2);w.setsampwidth(2);w.setframerate(48000);w.writeframes(bytes(48000*2*2))
print(json.dumps({'stage':'source_built','scenes':58,'chapters':11,'words':meta['estimated_spoken_words'],'deck_slides':18,'speech_generated':False,'mp4_rendered':False}))
