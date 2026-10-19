"""Narration-timed Chinese SVG scenes, captions, editable deck and chapter hosts."""
from pathlib import Path
import html,json,math,re,shutil,os
from .common import read,write,sha,digest
from .mandarin_art import build,PAPER,INK,WHITE,ACCENT,SOFT,GRAY

def esc(x):return html.escape(str(x))

def wrap(text,font,width):
    lines=[];line=''
    for char in text:
        if char=='\n':lines.append(line);line='';continue
        if line and font.getlength(line+char)>width:
            # Keep closing punctuation with its preceding phrase.
            if char in '，。！？；：、）》」』％':line+=char;continue
            lines.append(line);line=char
        else:line+=char
    if line:lines.append(line)
    return lines or ['']

def font(p,size,bold=False):
    from PIL import ImageFont
    return ImageFont.truetype(str(p.root/p.profile['fonts']['bold' if bold else 'regular']),size)

def text_svg(p,o):
    lines=wrap(o['text'],font(p,o['size'],o.get('bold',False)),o['w'])
    return f'<text x="{o["x"]}" y="{o["y"]+o["size"]}" font-size="{o["size"]}" font-weight="{700 if o.get("bold") else 400}" fill="{o.get("color",INK)}" text-anchor="{o.get("anchor","middle")}">'+''.join(f'<tspan x="{o["x"]}" dy="{0 if i==0 else o["size"]*1.25}">{esc(t)}</tspan>' for i,t in enumerate(lines))+'</text>'

def object_svg(p,o,hidden=False):
    k=o['kind']
    if k=='rect':body=f'<rect x="{o["x"]}" y="{o["y"]}" width="{o["w"]}" height="{o["h"]}" rx="{o["rx"]}" fill="{o["fill"]}" stroke="{o["stroke"]}" stroke-width="4"/>'
    elif k=='ellipse':body=f'<ellipse cx="{o["x"]+o["w"]/2}" cy="{o["y"]+o["h"]/2}" rx="{o["w"]/2}" ry="{o["h"]/2}" fill="{o["fill"]}" stroke="{o["stroke"]}" stroke-width="4"/>'
    elif k=='text':body=text_svg(p,o)
    elif k=='line':
        x1,y1,x2,y2=[o[n] for n in ['x1','y1','x2','y2']]
        d=f'M{x1} {y1} L{x2} {y2}'
        if o.get('arrow'):
            a=math.atan2(y2-y1,x2-x1);a1=a+.5;a2=a-.5
            d+=f' M{x2-18*math.cos(a1)} {y2-18*math.sin(a1)} L{x2} {y2} L{x2-18*math.cos(a2)} {y2-18*math.sin(a2)}'
        body=f'<path d="{d}" fill="none" stroke="{o["color"]}" stroke-width="4" stroke-linecap="round" stroke-linejoin="round"'+(' stroke-dasharray="9 9"' if o.get('dash') else '')+'/>'
    else:raise ValueError(k)
    return f'<g id="{o["id"]}"'+(' style="opacity:0"' if hidden and k!='line' else '')+' class="'+('hw-arrow' if k=='line' and not o.get('dash') else 'native-object')+'">'+body+'</g>'

def tally(p,sid,pose='point'):
    svg=(p.root/'assets/mascot'/f'tally_{pose}.svg').read_text()
    inner=re.sub(r'^.*?<svg[^>]*>|</svg>\s*$','',svg,flags=re.S)
    ids=re.findall(r'id="([^"]+)"',inner)
    for ident in ids:
        inner=inner.replace(f'id="{ident}"',f'id="{sid}-tally-{ident}"').replace(f'#{ident}',f'#{sid}-tally-{ident}')
    return f'<g id="{sid}-tally" transform="translate(1620 640) scale(.74)">{inner}</g>'

def css(p):return '''@font-face{font-family:EditorialTC;src:url(assets/fonts/NotoSansTC-Regular.otf)}@font-face{font-family:EditorialTC;src:url(assets/fonts/NotoSansTC-Bold.otf);font-weight:700}@font-face{font-family:Technical;src:url(assets/fonts/DejaVuSansMono.ttf)}*{box-sizing:border-box}html,body{margin:0}#root{position:relative;width:100%;height:100%;overflow:hidden;color:#292421;background:#F0EBE4;font-family:EditorialTC,sans-serif}.art{position:absolute;inset:0;width:1920px;height:1080px;font-family:EditorialTC,sans-serif}.chapter{position:absolute;left:145px;top:78px;font-size:27px;line-height:1.4}.heading{position:absolute;left:145px;top:120px;width:1620px;font-size:62px;font-weight:700;line-height:1.25;margin:0}.teaching{position:absolute;left:175px;top:870px;width:1390px;text-align:center;font-size:35px;font-weight:700;line-height:1.25}.caption{position:absolute;left:180px;top:942px;width:1560px;min-height:58px;max-height:112px;padding:4px 18px;text-align:center;font-size:38px;line-height:1.27;font-family:EditorialTC,sans-serif;background:#F0EBE4;white-space:normal;line-break:strict;overflow-wrap:normal}.question{position:absolute;left:210px;top:270px;width:1300px;font-size:42px;font-weight:700;line-height:1.25}.question-title{position:absolute;left:210px;top:220px;font-size:30px}.choices{position:absolute;left:210px;width:1300px;display:flex;flex-direction:column;gap:12px}.option{padding:12px 18px;border:2px solid #B8B0A9;background:#FFFAF3;font-size:34px;line-height:1.25}.answer-mark{position:absolute;left:210px;top:865px;width:1300px;text-align:center;font-size:34px;font-weight:700}.quiz-clock{position:absolute;left:1620px;top:320px;width:115px;height:115px;border:4px solid #292421;border-radius:50%;background:#F2D4B7;display:flex;align-items:center;justify-content:center;font-size:48px}.paper-lines{stroke:#DCD3C8;stroke-width:2;fill:none}'''

def caption_rows(p,sid,audio):
    result=[]
    for segment in audio['segments']:
        if segment['kind']!='spoken':continue
        alignment=read(p.root/'data/forced-alignment'/f'{sid}-{segment["sentence"]:03}.json')
        assert alignment['audio_sha256']==segment['sha256']
        words=alignment['words'];groups=[];group=[]
        for word in words:
            trial=''.join(x['text'] for x in group+[word])
            if group and (len(wrap(trial,font(p,38),1450))>2 or sum(1 if ord(c)>255 else .55 for c in trial)>43):
                groups.append(group);group=[]
            group.append(word)
            if word['text'].rstrip().endswith(('。','？','！','；')) and len(''.join(x['text'] for x in group))>=12:groups.append(group);group=[]
        if group:groups.append(group)
        if len(groups)>1 and len(''.join(w['text'] for w in groups[-1]).strip())<8:
            combined=groups[-2]+groups[-1]
            if len(wrap(''.join(w['text'] for w in combined),font(p,38),1450))<=2:
                groups[-2:]=[combined]
        for group in groups:
            text=''.join(x['text'] for x in group);start=segment['start_sample']/48000+group[0]['start'];end=segment['start_sample']/48000+group[-1]['end']
            assert end>start and len(wrap(text,font(p,38),1450))<=2
            result.append({'text':text,'start':start,'end':end,'source_sentence':segment['sentence'],'timing_source':'provider_word_boundary'})
        assert ''.join(x['text'] for x in result if x['source_sentence']==segment['sentence'])==segment['canonical']
    return result

def background():
    lines=''.join(f'<path d="M110 {y} Q930 {y+3} 1810 {y}"/>' for y in [225,920])
    return f'<svg class="art" viewBox="0 0 1920 1080" xmlns="http://www.w3.org/2000/svg"><path class="paper-lines" d="M112 65 Q103 540 115 1015 M1810 65 Q1818 540 1805 1015"/><g class="paper-lines">{lines}</g></svg>'

def scene_body(p,scene,audio):
    sid=scene['id'];parts=[background()];js=[];objects=build(scene)
    spoken=[r for r in audio['segments'] if r['kind']=='spoken'];phases=[r['start_sample']/48000 for r in spoken]
    def show(ident,start,end=None):
        js.append(f'gsap.set("#{ident}",{{opacity:0}});tl.set("#{ident}",{{opacity:1}},{start:.8f});')
        if end is not None:js.append(f'tl.set("#{ident}",{{opacity:0}},{end:.8f});')
    if scene['kind']=='quiz':
        q=next(q for q in p.questions if q['scene_id']==sid);hold=next(x for x in audio['segments'] if x['kind']=='silence');answer=hold['end_sample']/48000
        prompt_height=max(2,len(wrap(q['prompt'],font(p,42,True),1300)))*42*1.35+16;top=270+prompt_height+25
        option_heights=[len(wrap(c['letter']+'．'+c['text'],font(p,34),1260))*34*1.25+28 for c in q['choices']]
        assert top+sum(option_heights)+36<850,(sid,top,option_heights)
        parts.append(f'<div class="question-title">{esc(q["level"])}練習・需要時可暫停</div><div class="question">{esc(q["prompt"])}</div><div class="choices" style="top:{top}px">'+''.join(f'<div class="option" id="{sid}-{c["letter"]}">{c["letter"]}．{esc(c["text"])}</div>' for c in q['choices'])+'</div>')
        parts.append(f'<div class="quiz-clock" id="{sid}-clock" data-quiz-scene="{sid}" data-pause-start="{hold["start_sample"]/48000}" data-pause-end="{answer}">1</div>')
        show(sid+'-clock',hold['start_sample']/48000,answer)
        parts.append(f'<div class="answer-mark" id="{sid}-answer">正確答案：{q["correct_letter"]}</div>');show(sid+'-answer',answer)
        js.append(f'tl.to("#{sid}-{q["correct_letter"]}",{{backgroundColor:"{SOFT}",borderColor:"{INK}",duration:.35}},{answer});')
        parts.append('<svg class="art" viewBox="0 0 1920 1080" xmlns="http://www.w3.org/2000/svg">'+tally(p,sid,'think')+'</svg>')
        write(p.root/'data'/f'{sid}.quiz-layout.json',{'choice_end':top+sum(option_heights)+36,'answer_start':answer})
    else:
        parts.append('<svg class="art" viewBox="0 0 1920 1080" xmlns="http://www.w3.org/2000/svg">'+''.join(object_svg(p,o,hidden=True) for o in objects)+tally(p,sid)+'</svg>')
        counters={}
        for o in objects:
            phase=min(o['phase'],len(phases)-1);i=counters.get(phase,0);counters[phase]=i+1
            at=phases[phase]+min(i*.03,.45);ident=o['id']
            if o['kind']=='line' and not o.get('dash'):
                js.append(f'hwArrowOn(tl,"#{ident}",{at:.8f},.7);')
            else:
                js.append(f'tl.fromTo("#{ident}",{{opacity:0,y:{12 if i%2 else -8}}},{{opacity:1,y:0,duration:.55,ease:"{["power3.out","power2.out","sine.out"][i%3]}",immediateRender:false}},{at:.8f});')
            if o.get('motion')=='bar':js.append(f'tl.fromTo("#{ident}",{{scaleX:0,transformOrigin:"0% 50%"}},{{scaleX:1,duration:1.2,immediateRender:false}},{at:.8f});')
            if o.get('motion')=='needle':js.append(f'tl.fromTo("#{ident}",{{rotation:0,svgOrigin:"820 690"}},{{rotation:12,svgOrigin:"820 690",duration:1.4,repeat:3,yoyo:true,ease:"sine.inOut",immediateRender:false}},{at+1:.8f});')
            if o.get('travel'):
                token=ident+'-token';start=at+.8;duration=min(2.4,max(.5,audio['duration']-start-1))
                parts.append(f'<svg class="art" viewBox="0 0 1920 1080" xmlns="http://www.w3.org/2000/svg"><g id="{token}"><rect x="-14" y="-9" width="28" height="18" rx="3" fill="{ACCENT}" stroke="{INK}" stroke-width="2"/></g></svg>')
                js.append(f'tl.set("#{token}",{{opacity:0}},0);tl.fromTo("#{token}",{{x:{o["x1"]},y:{o["y1"]},opacity:1}},{{x:{o["x2"]},y:{o["y2"]},opacity:1,duration:{duration},ease:"none",immediateRender:false}},{start});tl.set("#{token}",{{opacity:0}},{start+duration});')
        parts.append(f'<div class="teaching" id="{sid}-takeaway">{esc(scene["takeaway"])}</div>');show(sid+'-takeaway',phases[-1])
    parts.append(f'<div class="chapter">{esc(scene["chapter_id"]+"　"+scene["chapter_title"])}</div><h1 class="heading">{esc(scene["title"])}</h1>')
    parts.append(f'<svg class="art" viewBox="0 0 1920 1080" xmlns="http://www.w3.org/2000/svg"><text x="1790" y="900" text-anchor="end" font-size="22" fill="{INK}">教學用模擬資料</text></svg>')
    captions=caption_rows(p,sid,audio)
    for i,cap in enumerate(captions):
        ident=f'{sid}-caption-{i}';parts.append(f'<div class="caption" id="{ident}">{esc(cap["text"])}</div>');show(ident,cap['start'],cap['end'])
    return ''.join(parts),'\n'.join(js),objects,captions

def build_scene(p,sid):
    scene=p.by_id[sid];audio=p.audio(sid);body,motion,objects,captions=scene_body(p,scene,audio)
    helper=(p.root/'assets/vendor/hw-arrow-helpers.js').read_text()
    script=helper+'\nconst tl=gsap.timeline({paused:true});\n'+motion+f'\nwindow.__timelines["{sid}"]=tl;'
    markup=f'<template><style>{css(p)}</style><div id="root" data-composition-id="{sid}" data-width="1920" data-height="1080" data-duration="{audio["duration"]:.9f}">{body}</div><script>{script}</script></template>'
    (p.root/'compositions'/f'{sid}.html').write_text(markup)
    write(p.root/'compositions'/f'{sid}.motion.json',{'scene':sid,'duration':audio['duration'],'rules':scene['motion_rules'],
      'entry':{'vector':'0 px/s','from_rest':True},'exit':{'vector':'0 px/s','still_moving':False},'objects':len(objects),'audio_sha256':sha(p.root/'assets/voice'/f'{sid}.wav')})
    write(p.root/'data'/f'{sid}.captions.json',captions);write(p.root/'data'/f'{sid}.geometry.json',objects)
    board=p.root/'boards/svg'/f'{sid}.svg';board.parent.mkdir(parents=True,exist_ok=True)
    header={'kind':'text','x':145,'y':125,'text':scene['title'],'size':62,'w':1610,'bold':True,'anchor':'start'}
    content=''.join(object_svg(p,o) for o in objects)
    if scene['kind']=='quiz':
        q=next(q for q in p.questions if q['scene_id']==sid)
        items=[dict(kind='text',x=210,y=260,text=q['prompt'],size=42,w=1300,bold=True,anchor='start')]
        for i,ch in enumerate(q['choices']):items.append(dict(kind='text',x=230,y=410+i*100,text=ch['letter']+'．'+ch['text'],size=34,w=1280,anchor='start'))
        content=''.join(text_svg(p,o) for o in items)
    board.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080"><style>text{{font-family:Noto Sans CJK TC,sans-serif}}</style><rect width="1920" height="1080" fill="{PAPER}"/>'+text_svg(p,header)+content+tally(p,sid)+'</svg>')

def host(p,scenes,name,folder):
    import soundfile as sf
    folder.mkdir(parents=True,exist_ok=True);assets=folder/'assets'
    for sub in ['fonts','vendor']:
        for source in (p.root/'assets'/sub).glob('*'):
            if not source.is_file():continue
            dest=assets/sub/source.name;dest.parent.mkdir(parents=True,exist_ok=True)
            if source.resolve()!=dest.resolve() and (not dest.exists() or sha(source)!=sha(dest)):shutil.copy2(source,dest)
    duration=0;slots=[];audio=assets/'audio'/f'{name}.wav';audio.parent.mkdir(parents=True,exist_ok=True)
    with sf.SoundFile(str(audio)+'.tmp.wav','w',samplerate=48000,channels=2,subtype='PCM_16') as out:
        for scene in scenes:
            sid=scene['id'];a=p.audio(sid);d=a['duration']
            slots.append(f'<div class="clip" id="slot-{sid}" data-composition-id="{sid}" data-composition-src="compositions/{sid}.html" data-start="{duration:.9f}" data-duration="{d:.9f}" data-track-index="1" data-width="1920" data-height="1080" style="position:absolute;inset:0"></div>')
            pcm,rate=sf.read(p.root/'assets/voice'/f'{sid}.wav',dtype='int16');assert rate==48000;out.write(pcm)
            if folder!=p.root:
                (folder/'compositions').mkdir(exist_ok=True)
                for suffix in ['html','motion.json']:shutil.copy2(p.root/'compositions'/f'{sid}.{suffix}',folder/'compositions'/f'{sid}.{suffix}')
            duration+=d
    Path(str(audio)+'.tmp.wav').replace(audio)
    slots.append(f'<audio class="clip" id="narration-{name}" src="assets/audio/{name}.wav" data-start="0" data-duration="{duration:.9f}" data-track-index="10"></audio>')
    markup=f'<!doctype html><html lang="zh-TW"><head><meta charset="utf-8"><script src="assets/vendor/gsap.min.js"></script><style>{css(p)}</style></head><body><div id="root" data-composition-id="{name}" data-width="1920" data-height="1080" data-duration="{duration:.9f}">'+''.join(slots)+f'</div><script>window.__timelines["{name}"]=gsap.timeline({{paused:true}});</script></body></html>'
    (folder/'index.html').write_text(markup)
    if folder!=p.root:shutil.copy2(p.root/'hyperframes.json',folder/'hyperframes.json')
    return duration

def assemble(p):
    duration=host(p,p.scenes,'master',p.root);offset=0;rows=[];silences=[];captions=[];chapters=[]
    for ch in p.chapters:
        subset=p.chapter_scenes(ch);d=sum(p.audio(s['id'])['duration'] for s in subset)
        chapters.append({'id':ch,'title':subset[0]['chapter_title'],'start':offset,'duration':d})
        for scene in subset:
            sid=scene['id'];a=p.audio(sid);rows.append({'scene':sid,'chapter':ch,'start':offset,'duration':a['duration'],'frames':round(a['duration']*30)})
            for seg in a['segments']:
                if seg['kind']=='silence':silences.append({'scene':sid,'start':offset+seg['start_sample']/48000,'end':offset+seg['end_sample']/48000,'start_sample':round(offset*48000)+seg['start_sample'],'end_sample':round(offset*48000)+seg['end_sample']})
            for cap in read(p.root/'data'/f'{sid}.captions.json'):captions.append({**cap,'scene':sid,'start':offset+cap['start'],'end':offset+cap['end']})
            offset+=a['duration']
    assert abs(offset-duration)<.000001
    write(p.root/'data/master-timing.json',{'duration':duration,'frames':round(duration*30),'scenes':rows,'chapters':chapters,'silences':silences})
    write(p.root/'data/master-captions.json',captions)
    def clock(t,sep):
        ms=round(t*1000);h,ms=divmod(ms,3600000);m,ms=divmod(ms,60000);s,ms=divmod(ms,1000)
        return f'{h:02}:{m:02}:{s:02}{sep}{ms:03}'
    for ext,sep in [('srt',','),('vtt','.')]:
        blocks=['WEBVTT\n'] if ext=='vtt' else []
        for i,cap in enumerate(captions,1):blocks.append(f'{i}\n{clock(cap["start"],sep)} --> {clock(cap["end"],sep)}\n{cap["text"]}\n')
        (p.output/f'captions.zh-TW.{ext}').write_text('\n'.join(blocks))
    return {'duration':duration,'frames':round(duration*30),'captions':len(captions)}

def deck(p):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    scale=914400/120
    def rgb(s):return RGBColor.from_string(s.lstrip('#'))
    def shape(slide,o):
        if o['kind'] in ['rect','ellipse']:
            kind=MSO_SHAPE.RECTANGLE if o['kind']=='rect' else MSO_SHAPE.OVAL
            s=slide.shapes.add_shape(kind,int(o['x']*scale),int(o['y']*scale),int(o['w']*scale),int(o['h']*scale))
            if o['fill']=='none':s.fill.background()
            else:s.fill.solid();s.fill.fore_color.rgb=rgb(o['fill'])
            s.line.color.rgb=rgb(o['stroke']);s.line.width=Pt(1.8)
        elif o['kind']=='line':
            s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,int(o['x1']*scale),int(o['y1']*scale),int(o['x2']*scale),int(o['y2']*scale));s.line.color.rgb=rgb(o['color']);s.line.width=Pt(1.8)
            if o.get('arrow'):
                arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');s.line._get_or_add_ln().append(arrow)
        elif o['kind']=='text':
            lines=wrap(o['text'],font(p,o['size'],o.get('bold',False)),o['w']);height=len(lines)*o['size']*1.35+8
            left=o['x']-o['w']/2 if o.get('anchor','middle')=='middle' else o['x']
            s=slide.shapes.add_textbox(int(left*scale),int(o['y']*scale),int(o['w']*scale),int(height*scale));tf=s.text_frame;tf.clear();tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.word_wrap=True
            for i,text in enumerate(lines):
                para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=text;para.alignment=PP_ALIGN.CENTER if o.get('anchor','middle')=='middle' else PP_ALIGN.LEFT
                para.font.name='Noto Sans CJK TC';para.font.size=Pt(o['size']*.6);para.font.bold=o.get('bold',False);para.font.color.rgb=rgb(o.get('color',INK));para.space_after=Pt(0)
                for run in para.runs:
                    rpr=run._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface','Noto Sans CJK TC');rpr.append(ea)
        else:raise ValueError(o['kind'])
    for scene in p.scenes:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(PAPER)
        for o in read(p.root/'data'/f'{scene["id"]}.geometry.json'):shape(slide,o)
        if scene['kind']=='quiz':
            q=next(q for q in p.questions if q['scene_id']==scene['id'])
            shape(slide,dict(kind='text',x=210,y=260,text=q['prompt'],size=42,w=1300,bold=True,anchor='start'))
            for i,ch in enumerate(q['choices']):shape(slide,dict(kind='text',x=230,y=410+i*100,text=ch['letter']+'．'+ch['text'],size=34,w=1280,anchor='start'))
        shape(slide,dict(kind='text',x=145,y=125,text=scene['title'],size=62,w=1610,bold=True,anchor='start'))
        shape(slide,dict(kind='text',x=145,y=80,text=scene['chapter_id']+'　'+scene['chapter_title'],size=27,w=1600,anchor='start'))
        shape(slide,dict(kind='text',x=850,y=875,text=scene['takeaway'],size=35,w=1400,bold=True))
        slide.shapes.add_picture(str(p.root/'assets/mascot/tally_point.png'),int(1620*scale),int(640*scale),width=int(200*scale))
        slide.notes_slide.notes_text_frame.text=scene['narration_zh_TW']
    target=p.output/'Digital_Everyday_Masterclass_zh-TW_EDITABLE.pptx';prs.save(target)
    return {'slides':len(prs.slides),'path':str(target),'native_editable_shapes':sum(len(s.shapes) for s in prs.slides)}
